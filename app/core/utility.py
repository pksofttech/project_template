"""Central Utility Functions: DataTables, Excel Export, SSE, Performance Timers."""

import asyncio
import csv
import functools
import inspect
import json
import re
import time
from collections.abc import Callable
from datetime import date, datetime, timedelta
from io import BytesIO, StringIO
from typing import Any
from zoneinfo import ZoneInfo

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core.models import MODEL_MAP
from app.stdio import print_debug, print_error, print_warning, time_now

# SSE Client Connections List
sse_clients: list[asyncio.Queue] = []


def time_logger(func: Callable) -> Callable:
    """Timer Decorator for tracking execution performance of async and sync functions."""
    is_async = inspect.iscoroutinefunction(func)

    if is_async:

        @functools.wraps(func)
        async def async_wrapper(*args: Any, **kwargs: Any) -> Any:
            t0 = time.perf_counter()
            try:
                return await func(*args, **kwargs)
            finally:
                ms = (time.perf_counter() - t0) * 1000
                print_debug(f"⏱️ [Async] {func.__name__} took: {ms:.2f} ms")

        return async_wrapper

    @functools.wraps(func)
    def sync_wrapper(*args: Any, **kwargs: Any) -> Any:
        t0 = time.perf_counter()
        try:
            return func(*args, **kwargs)
        finally:
            ms = (time.perf_counter() - t0) * 1000
            print_debug(f"⏱️ [Sync] {func.__name__} took: {ms:.2f} ms")

    return sync_wrapper


def broadcast_sse(payload: dict):
    """Broadcast an SSE event dictionary to all active connected clients."""

    async def _send_all():
        dead_clients = []
        for q in sse_clients:
            try:
                q.put_nowait(payload)
            except asyncio.QueueFull:
                dead_clients.append(q)

        for q in dead_clients:
            if q in sse_clients:
                sse_clients.remove(q)

    asyncio.create_task(_send_all())


def get_datatable_select(params: dict) -> dict:
    """
    Parse DataTables server-side request parameters including pagination,
    column-level searches, ordering, date ranges, and custom filters.
    """
    try:
        skip = int(params.get("start", "0"))
    except (ValueError, TypeError):
        skip = 0

    try:
        limit = int(params.get("length", "10"))
        if limit < 0:
            limit = None
    except (ValueError, TypeError):
        limit = 10

    target_table = (params.get("table") or "").strip()

    list_datas = []
    for k in params:
        match = re.search(r"^columns\[.*\]\[data\]", k)
        if match:
            data_val = (params[k] or "").strip()
            index_cols = k.split("[")[1].split("]")[0]
            search = params.get(f"columns[{index_cols}][search][value]", "")
            name = params.get(f"columns[{index_cols}][name]", "")

            if "." in data_val:
                _table, _col = data_val.split(".", 1)
                list_datas.append({"table": _table, "col": _col, "search": search, "name": name or _col})
            elif "_" in data_val:
                matched_model = False
                for table_name in MODEL_MAP:
                    if data_val.lower().startswith(f"{table_name.lower()}_"):
                        _table = table_name
                        _col = data_val[len(table_name) + 1 :]
                        list_datas.append({"table": _table, "col": _col, "search": search, "name": data_val})
                        matched_model = True
                        break
                if not matched_model and target_table:
                    list_datas.append({"table": target_table, "col": data_val, "search": search, "name": name or data_val})
            elif target_table and data_val:
                list_datas.append({"table": target_table, "col": data_val, "search": search, "name": name or data_val})
            elif name and "." in name:
                _table, _col = name.split(".", 1)
                list_datas.append({"table": _table, "col": _col, "search": search, "name": data_val or _col})

    order_by = None
    order_idx = params.get("order[0][column]")
    if order_idx:
        order_data = (params.get(f"columns[{order_idx}][data]") or "").strip()
        order_name = (params.get(f"columns[{order_idx}][name]") or "").strip()
        _table, _col = None, None
        if "." in order_data:
            _table, _col = order_data.split(".", 1)
        elif target_table and order_data:
            _table, _col = target_table, order_data
        elif "." in order_name:
            _table, _col = order_name.split(".", 1)
        elif target_table and order_name:
            _table, _col = target_table, order_name

        if _table and _col:
            order_by = {
                "table": _table,
                "col": _col,
                "dir": params.get("order[0][dir]", "asc"),
            }

    date_range = params.get("date_range")
    if date_range:
        date_range = date_range.replace("/", "-")
        try:
            if len(date_range) == 10:
                d_start = datetime.strptime(date_range, "%Y-%m-%d")
                d_end = d_start + timedelta(days=1)
            else:
                date_format = "%Y-%m-%d %H:%M"
                start_str, end_str = date_range.split(" - ")
                d_start = datetime.strptime(start_str, date_format)
                d_end = datetime.strptime(end_str, date_format)

            d_start = d_start.replace(tzinfo=ZoneInfo("Asia/Bangkok")).isoformat()
            d_end = d_end.replace(tzinfo=ZoneInfo("Asia/Bangkok")).isoformat()
            date_range = (d_start, d_end)
        except ValueError as err:
            print_error(f"Date parsing error: {err}")
            date_range = None

    filter_dict = {}
    _filter = params.get("filter")
    if _filter:
        try:
            filter_dict = json.loads(_filter)
        except (json.JSONDecodeError, TypeError) as err:
            print_debug(f"Filter decode error: {err}")

    return {
        "table": params.get("table"),
        "list_datas": list_datas,
        "order_by": order_by,
        "skip": skip,
        "limit": limit,
        "filter": filter_dict,
        "date_range": date_range,
        "search": params.get("search[value]", "").strip(),
    }


async def export_excel_response(
    rows: list,
    filename_prefix: str = "export",
    sheet_title: str = "Data Export",
    export_type: str = "excel",
) -> StreamingResponse:
    """
    Dynamically converts a list of mappings/dicts into an Excel (.xlsx) or CSV file
    with automatic column formatting and UTF-8 BOM support.
    """

    def flatten_row(row):
        flat = {}
        for k, val in row.items():
            if val is None:
                flat[k] = ""
            elif hasattr(val, "__table__"):
                for col in val.__table__.columns:
                    flat[f"{k}.{col.name}"] = getattr(val, col.name)
            elif isinstance(val, dict):
                for sub_k, sub_v in val.items():
                    flat[f"{k}.{sub_k}"] = sub_v
            elif isinstance(val, list):
                flat[k] = ", ".join(str(item) for item in val)
            elif isinstance(val, (datetime, date)):
                flat[k] = val.strftime("%Y-%m-%d %H:%M:%S") if isinstance(val, datetime) else val.strftime("%Y-%m-%d")
            else:
                flat[k] = val
        return flat

    flat_rows = [flatten_row(r) for r in rows] if rows else []
    keys = list(dict.fromkeys(k for r in flat_rows for k in r)) if flat_rows else []
    headers = [k.split(".")[-1].replace("_", " ").title() for k in keys] if keys else ["No Records"]

    _now = time_now().strftime("%Y%m%d_%H%M%S")

    if export_type == "csv":
        csv_buffer = StringIO()
        csv_buffer.write("\ufeff")  # UTF-8 BOM for Thai language in Excel
        writer = csv.writer(csv_buffer)
        writer.writerow(headers)
        for row in flat_rows:
            writer.writerow([row.get(k, "") for k in keys])

        csv_buffer.seek(0)
        return StreamingResponse(
            iter([csv_buffer.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{filename_prefix}_{_now}.csv"'},
        )

    # Excel Export (.xlsx) via openpyxl
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:30]

    header_font = Font(name="Sarabun", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    cell_font = Font(name="Sarabun", size=10)
    thin_border = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB"),
    )

    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx, row in enumerate(flat_rows, start=2):
        row_values = [row.get(k, "") for k in keys]
        ws.append(row_values)
        for c_idx in range(1, len(row_values) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = cell_font
            cell.border = thin_border

    # Auto-fit column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)

    return StreamingResponse(
        excel_buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename_prefix}_{_now}.xlsx"'},
    )
